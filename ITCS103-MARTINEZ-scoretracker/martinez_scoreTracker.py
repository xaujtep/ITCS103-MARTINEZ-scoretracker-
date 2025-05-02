import tkinter as tk
from tkinter import messagebox
import openpyxl
import os

filename = "student_scores.xlsx"

def create_workbook():
    if not os.path.exists(filename):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Scores"
        ws.append(["Name", "Score", "Status"])
        wb.save(filename)

def student_score(name, score):
    wb = openpyxl.load_workbook(filename)
    ws = wb["Scores"]
    if score >= 30:
        status = "Pass"
    else:
        status = "Fail"
    ws.append([name, score, status])
    wb.save(filename)

def saving_student_record():
    name = name_entry.get()
    score_text = score_entry.get()

    if name == "" or score_text == "":
        messagebox.showerror("Error", "Please enter both name and score.")
        return

    if score_text.isdigit():
        score = float(score_text)
        student_score(name, score)
        messagebox.showinfo("Saved", "Record saved.")
        name_entry.delete(0, tk.END)
        score_entry.delete(0, tk.END)
    else:
        messagebox.showerror("Error", "Score must be a number.")

def records():
    wb = openpyxl.load_workbook(filename)
    ws = wb["Scores"]
    rec_window = tk.Toplevel(window)
    rec_window.title("Student Records")
    for row in ws.iter_rows(min_row=2, values_only=True):
        text = f"{row[0]} - {row[1]} - {row[2]}"
        label = tk.Label(rec_window, text=text)
        label.pack()

create_workbook()

window = tk.Tk()
window.geometry("200x200")
window.title("Student Score Tracker")

tk.Label(window, text="Name:").grid(row=0, column=0, padx=5, pady=5)
tk.Label(window, text="Score:").grid(row=1, column=0, padx=5, pady=5)

name_entry = tk.Entry(window)
score_entry = tk.Entry(window)

name_entry.grid(row=0, column=1, padx=5, pady=5)
score_entry.grid(row=1, column=1, padx=5, pady=5)

tk.Button(window, text="Save", command=saving_student_record, bg="lightgreen").grid(row=2, column=0, columnspan=2, pady=10)
tk.Button(window, text="Show Records", command=records, bg="lightblue").grid(row=3, column=0, columnspan=2, pady=5)

window.mainloop()
